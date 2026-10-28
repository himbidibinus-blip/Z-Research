"""v0.9.3 Data Consistency & Research Integrity: API, concurrency, source fidelity."""
import concurrent.futures
import hashlib
import io
import json
import multiprocessing
import os
import subprocess
import zipfile
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from app import main, ai_features, v05, v09
from app.citations import render_document, _author_people


def owner():
    client = TestClient(main.app)
    response = client.post('/api/auth/setup', json={'username':'qa093','password':'strong-password-093'})
    assert response.status_code == 200, response.text
    return client


def import_dataset(client, rows=20):
    source = ('X,Y\n' + ''.join(f'{index},{index * 2}\n' for index in range(rows))).encode()
    response = client.post('/api/datasets', files={'file': ('source.csv', source)})
    assert response.status_code == 200, response.text
    return response.json()['id'], source


def edit(client, did, row, value):
    return client.post(f'/api/v05/datasets/{did}/edit-cell', json={'row':row,'column':'X','value':value})


def audit_count(did):
    with ai_features.db() as db:
        return db.execute('SELECT COUNT(*) FROM dataset_edits WHERE dataset_id=?', (did,)).fetchone()[0]


def test_parallel_24_distinct_edits_all_persist_and_have_one_audit_event_each():
    client=owner()
    did, source=import_dataset(client,30)
    with concurrent.futures.ThreadPoolExecutor(max_workers=12) as pool:
        outcomes=list(pool.map(lambda row: edit(client,did,row,1000+row),range(24)))
    assert all(o.status_code==200 for o in outcomes),[(o.status_code,o.text) for o in outcomes]
    data,_=main.read_df(did)
    assert data['X'].iloc[:24].tolist()==list(range(1000,1024))
    assert audit_count(did)==24
    assert (main.DATA/(did+'.original.csv')).read_bytes()==source
    assert hashlib.sha256((main.DATA/'raw_uploads'/(did+'.csv')).read_bytes()).hexdigest()==hashlib.sha256(source).hexdigest()
    assert not list(main.DATA.glob('*.editing.csv'))
    assert not list(main.DATA.glob('*.rollback.csv'))


def _fork_edit(args):
    did,row,value,token=args
    from fastapi.testclient import TestClient
    from app.main import app
    client=TestClient(app)
    client.cookies.set('qr_session',token)
    outcome=client.post(f'/api/v05/datasets/{did}/edit-cell',json={'row':row,'column':'X','value':value})
    return outcome.status_code, outcome.text


@pytest.mark.skipif(os.name!='posix' or 'fork' not in multiprocessing.get_all_start_methods(),reason='POSIX process fork smoke; Windows filelock uses OS-native Windows lock')
def test_six_independent_server_processes_preserve_every_edit():
    client=owner()
    did,_=import_dataset(client,12)
    token=client.cookies['qr_session']
    with multiprocessing.get_context('fork').Pool(6) as pool:
        outcomes=pool.map(_fork_edit,[(did,i,200+i,token) for i in range(6)])
    assert all(code==200 for code,_ in outcomes),outcomes
    data,_=main.read_df(did)
    assert data['X'].iloc[:6].tolist()==list(range(200,206))
    assert audit_count(did)==6



def make_corrupt_workbook():
    blob=io.BytesIO()
    with zipfile.ZipFile(blob,'w') as archive:
        archive.writestr('xl/worksheets/sheet1.xml','<worksheet>')
    return blob

@pytest.mark.parametrize('raw',[
    b'not an xlsx file',
    b'PK\x03\x04corrupt',
    b'',
    b'\x00\x01\x02\xff',
    make_corrupt_workbook().getvalue(),
])
def test_malformed_xlsx_rejected_without_partial_data(raw):
    client=owner()
    result=client.post('/api/datasets',files={'file':('corrupt.xlsx',raw)})
    assert result.status_code==422, (result.status_code,result.text)
    assert client.get('/api/datasets').json()==[]
    assert not list(main.DATA.glob('*.csv'))
    assert not list((main.DATA/'raw_uploads').glob('*')) if (main.DATA/'raw_uploads').exists() else True



def test_mutation_and_audit_log_not_committed_when_atomic_write_fails(monkeypatch):
    client=owner()
    did,original=import_dataset(client,3)
    def refuse(*args,**kwargs):raise OSError('simulated write failure')
    monkeypatch.setattr(v05,'_atomic_csv_replace',refuse)
    with TestClient(main.app,raise_server_exceptions=False) as isolated:
        isolated.cookies.set('qr_session',client.cookies['qr_session'])
        result=edit(isolated,did,0,100)
    assert result.status_code==500
    assert (main.DATA/(did+'.csv')).read_bytes()==original
    assert audit_count(did)==0


def test_old_values_in_edit_audit_log_read_from_latest_locked_version():
    client=owner();did,_=import_dataset(client,3)
    assert edit(client,did,0,50).status_code==200
    assert edit(client,did,0,60).status_code==200
    with ai_features.db() as db:
        rows=db.execute('SELECT old_value,new_value FROM dataset_edits WHERE dataset_id=? ORDER BY created_at,id',(did,)).fetchall()
    assert len(rows)==2
    assert rows[0]['old_value']=='0' and rows[0]['new_value']=='50.0'
    assert rows[1]['old_value']=='50' and rows[1]['new_value']=='60.0'


def test_manual_author_year_review_without_selected_refs_and_mixed_linked():
    client=owner(); did,_=import_dataset(client,5)
    pid=client.post('/api/v09/projects',json={'name':'Manual citation alert','dataset_id':did}).json()['id']
    payload={'body':'Prior work (Smith, 2025) showed patterns; Doe (2024) disagreed.','reference_ids':[],'style':'apa'}
    save=client.post(f'/api/v09/projects/{pid}/chapters/2/save',json=payload)
    assert save.status_code==200,save.text
    report=client.get(f'/api/v09/projects/{pid}/audit').json()
    assert any(x['severity']=='review' and 'Chapter 2' in x['message'] and 'Smith' in x['message'] for x in report['alerts'])
    chapter=client.get(f'/api/v09/projects/{pid}').json()['chapters']['2']
    assert chapter['status']=='current'  # source provenance separate from bibliographic review
    assert chapter['citation_integrity']=='unverified_plain_text'
    assert len(chapter['unlinked_author_year_candidates'])==2
    assert not v09.likely_unlinked_citations('The response increased in 2025 over 2024 (N=120).')
    # Even a mix of Cite+ and unlinked author-year requires review.
    ref=client.post('/api/v05/references/manual',json={'title':'A study','authors':'Smith, John','year':2025}).json()['id']
    mixed={'body':f'Linked [@ref:{ref}], but also (Doe, 2024).','reference_ids':[ref],'style':'apa'}
    assert client.post(f'/api/v09/projects/{pid}/chapters/2/save',json=mixed).status_code==200
    report=client.get(f'/api/v09/projects/{pid}/audit').json()
    assert any('Doe' in x['message'] for x in report['alerts'])


CITATION_CORPUS=[
    'Smith, John AND Doe, Jane',
    'Smith, John AnD Doe, Jane',
    'Smith, John and Doe, Jane',
    'Smith, John; Doe, Jane',
    'Smith, John, Doe, Jane',
    'John Smith, Jane Doe',
    'Smith, John',
    'John Smith and Jane Doe',
    '{World Health Organization}',
    'World Health Organization',
    'Anna van der Berg',
    'O’Neill, Catherine',
    'de la Cruz, Ana',
    '',
]


def test_python_and_javascript_citation_parity_through_full_rendered_document():
    samples=[]
    for i,authors in enumerate(CITATION_CORPUS):
        ref={'id':f'ref{i}','title':f'Study {i}', 'authors':authors,'year':2025,'venue':'Journal'}
        for style in ('apa','ieee'):
            body=f'This article [@ref:ref{i}]. [@bibliography]'
            samples.append({'ref':ref,'style':style,'body':body})
    payload=json.dumps(samples)
    js="""const Citation=require('./app/static/citation09.js');
const cases=JSON.parse(process.argv[1]);
process.stdout.write(JSON.stringify(cases.map(({ref,style,body})=>({
people:Citation.parseAuthors(ref.authors).map(a=>({family:a.family,given:a.given,organization:!!a.organization})),
body:Citation.renderDocument(body,style,[ref]).text}))));"""
    node=subprocess.run(['node','-e',js,payload],text=True,capture_output=True,check=True)
    results=json.loads(node.stdout)
    for item,result in zip(samples,results):
        expected=render_document(item['body'],item['style'],[item['ref']])['body']
        actual=result['body']
        assert expected==actual, f"Parity failed for {item['ref']['authors']!r}, {item['style']}: {expected!r} != {actual!r}"
        assert _author_people(item['ref']['authors'])==result['people']


def test_apa_word_export_matches_txt_and_js_preview_uppercase_and():
    client=owner();did,_=import_dataset(client,3)
    pid=client.post('/api/v09/projects',json={'name':'Citation Word/TXT','dataset_id':did}).json()['id']
    r=client.post('/api/v05/references/manual',json={'title':'Study','authors':'Smith, John AND Doe, Jane','year':2025,'venue':'Journal'})
    assert r.status_code==200,r.text
    rid=r.json()['id']
    draft=f'Results [@ref:{rid}]. [@bibliography]'
    resp=client.post(f'/api/v09/projects/{pid}/chapters/2/save',json={'body':draft,'reference_ids':[rid],'style':'apa'})
    assert resp.status_code==200,resp.text
    exported=client.get(f'/api/v09/projects/{pid}/chapters/2/docx')
    assert exported.status_code==200
    paragraphs='\n'.join(p.text for p in Document(io.BytesIO(exported.content)).paragraphs)
    assert '(Smith & Doe, 2025)' in paragraphs
    assert 'Smith, J., Doe, J.' in paragraphs
    ref={'id':rid,'title':'Study','authors':'Smith, John AND Doe, Jane','year':2025,'venue':'Journal','doi':'','url':''}
    txt=render_document(draft,'apa',[ref])['body']
    assert '(Smith & Doe, 2025)' in txt and 'Smith, J., Doe, J.' in txt
    js="const C=require('./app/static/citation09.js');let refs=JSON.parse(process.argv[1]);console.log(C.renderDocument(process.argv[2],'apa',refs).text)"
    browser=subprocess.run(['node','-e',js,json.dumps([ref]),draft],text=True,capture_output=True,check=True).stdout.rstrip('\n')
    assert browser==txt
