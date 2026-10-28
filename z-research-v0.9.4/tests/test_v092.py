"""v0.9.2 remediation release gates, including negative inputs and restorative data integrity."""
import hashlib
import io
import re
from pathlib import Path
from fastapi.testclient import TestClient
import pandas as pd
import pytest
from app.main import app
from app.citations import cited_ids, validate_linked_sources, render_document
from app import ai_features, v09

SAMPLE=Path('sample/consumer_survey_demo.csv').read_bytes()

def owner():
    cl=TestClient(app)
    assert cl.post('/api/auth/setup',json={'username':'test092','password':'long-password-092'}).status_code==200
    return cl

def project(cl):
    r=cl.post('/api/datasets',files={'file':('sample.csv',SAMPLE)})
    assert r.status_code==200, r.text
    p=cl.post('/api/v09/projects',json={'name':'Version restore checks','dataset_id':r.json()['id']})
    assert p.status_code==200,p.text
    return p.json()['id']

def ref(cl, title, author, year):
    r=cl.post('/api/v05/references/manual',json={'title':title,'authors':author,'year':year,'doi':'10.0012/'+title.replace(' ','-')})
    assert r.status_code==200,r.text
    return r.json()['id']

@pytest.mark.parametrize('source',[b'age,age,target\n1,2,3\n', b'age, age,target\n1,2,3\n',b'\xef\xbb\xbfAge,age,target\n1,2,3\n',b'age,\tage\t,target\n1,2,3\n',b'age,AGE,target\n1,2,3\n', b'\xef\xbb\xbfage,\xef\xbb\xbfage,target\n1,2,3\n'])
def test_header_collisions_are_rejected_before_pandas_and_file_storage(source):
    from app import main
    cl=owner()
    res=cl.post('/api/datasets',files={'file':('duplicates.csv',source,'text/csv')})
    assert res.status_code==422,res.text
    assert cl.get('/api/datasets').json()==[]
    assert not list((main.DATA/'raw_uploads').glob('*')) if (main.DATA/'raw_uploads').exists() else True

@pytest.mark.parametrize('source',[b'age,,target\n1,2,3\n',b'\xef\xbb\xbf,other\n1,2\n',b''])
def test_blank_header_rejected(source):
    cl=owner()
    assert cl.post('/api/datasets',files={'file':('blank.csv',source)}).status_code==422

@pytest.mark.parametrize('source',[b'\xef\xbb\xbfage,target\n1,3\n2,4\n',b'"last, first",target\n1,3\n2,4\n',b'age,age2,target\n1,2,3\n2,3,4\n'])
def test_valid_headers_preserved(source):
    cl=owner()
    res=cl.post('/api/datasets',files={'file':('valid.csv',source)})
    assert res.status_code==200,res.text
    assert cl.get('/api/datasets/'+res.json()['id']+'/provenance').json()['sha256']==hashlib.sha256(source).hexdigest()

def test_xlsx_duplicate_header_detected_before_pandas():
    cl=owner()
    df=pd.DataFrame([[1,2],[3,4]],columns=['age','age'])
    file=io.BytesIO();df.to_excel(file,index=False)
    r=cl.post('/api/datasets',files={'file':('duplicate.xlsx',file.getvalue())})
    assert r.status_code==422, r.text

def test_restore_replays_exact_body_source_selection_style_and_keeps_history():
    cl=owner();pid=project(cl)
    smith=ref(cl,'A study','Smith, John',2025)
    doe=ref(cl,'B study','Doe, Jane',2026)
    url=f'/api/v09/projects/{pid}/chapters/2'
    first=cl.post(url+'/save',json={'body':'First '+f'[@ref:{doe}]'+ ' [@bibliography]','reference_ids':[doe],'style':'ieee'}).json()
    second=cl.post(url+'/save',json={'body':'Second '+f'[@ref:{smith}]','reference_ids':[smith],'style':'apa'}).json()
    assert first['id']!=second['id']
    resp=cl.post(url+f"/versions/{first['id']}/restore",json={})
    assert resp.status_code==200,resp.text
    data=resp.json()
    assert data['body']=='First '+f'[@ref:{doe}]'+ ' [@bibliography]'
    assert data['reference_ids']==[doe] and data['citation_style']=='ieee'
    revisions=cl.get(url).json()
    assert len(revisions)==3
    assert revisions[0]['id']==data['id'] and revisions[0]['body']==data['body']
    assert revisions[0]['citation_style']=='ieee' and revisions[0]['reference_ids']==[doe]
    doc=cl.get(url+'/docx')
    from docx import Document
    full='\n'.join(x.text for x in Document(io.BytesIO(doc.content)).paragraphs)
    assert '[1]' in full and 'Doe' in full and '[@ref:' not in full
    assert full.count('References')==1


def test_restore_fails_if_historical_ref_was_deleted_and_does_not_create_revision():
    cl=owner();pid=project(cl)
    rid=ref(cl,'Vanished','Doe, Jane',2024)
    url=f'/api/v09/projects/{pid}/chapters/2'
    old=cl.post(url+'/save',json={'body':f'Case [@ref:{rid}]','reference_ids':[rid],'style':'ieee'}).json()['id']
    cl.post(url+'/save',json={'body':'Latest text','reference_ids':[],'style':'apa'})
    before=len(cl.get(url).json())
    assert cl.delete(f'/api/v05/references/{rid}').status_code==200
    result=cl.post(url+f'/versions/{old}/restore',json={})
    assert result.status_code==422,result.text
    assert len(cl.get(url).json())==before
    assert cl.get(url).json()[0]['body']=='Latest text'


def test_restore_stale_revision_not_misrepresented_as_current():
    cl=owner();pid=project(cl)
    rid=ref(cl,'Source','Smith, John',2025)
    url=f'/api/v09/projects/{pid}/chapters/1'
    old=cl.post(url+'/save',json={'body':f'Historical [@ref:{rid}]','reference_ids':[rid],'style':'apa'}).json()['id']
    with ai_features.db() as c:c.execute('UPDATE research_references SET title=? WHERE id=?',('Changed reference',rid))
    out=cl.post(url+f'/versions/{old}/restore',json={})
    assert out.status_code==200,out.text
    assert out.json()['status']=='stale'
    assert out.json()['warning']
    assert cl.get(url).json()[0]['status']=='stale'


def test_restore_rejects_foreign_chapter_id_and_anonymous_users():
    cl=owner();pid=project(cl)
    old=cl.post(f'/api/v09/projects/{pid}/chapters/2/save',json={'body':'Chapter 2'}).json()['id']
    assert cl.post(f'/api/v09/projects/{pid}/chapters/1/versions/{old}/restore',json={}).status_code==404
    anon=TestClient(app)
    assert anon.post(f'/api/v09/projects/{pid}/chapters/2/versions/{old}/restore',json={}).status_code==401
    assert cl.post(f'/api/v09/projects/{pid}/chapters/2/versions/{old}/restore',json={},headers={'Origin':'https://hostile.example'}).status_code==403


def test_ieee_numbering_is_first_occurrence_and_resynchronized_after_source_edits():
    refs=[{'id':'smith','title':'A','authors':'Smith, John','year':2025}, {'id':'doe','title':'B','authors':'Doe, Jane','year':2026}]
    body='Doe [@ref:doe], then Smith [@ref:smith] and Doe again [@ref:doe]. [@bibliography]'
    first=render_document(body,'ieee',refs)
    assert 'Doe [1], then Smith [2] and Doe again [1]' in first['body']
    assert first['body'].index('[1] J. Doe')<first['body'].index('[2] J. Smith')
    edited=render_document('Doe only [@ref:doe]. [@bibliography]','ieee',refs)
    assert 'Doe only [1]' in edited['body']
    assert '[2]' not in edited['body']
    assert 'Smith' not in edited['bibliography']
    assert render_document(body,'apa',refs)['body'].count('(Doe, 2026)')==2
    assert render_document(body,'apa',refs)['body'].count('(Smith, 2025)')==1


def test_server_rejects_orphan_tokens_and_missing_sources_not_reported_valid():
    cl=owner();pid=project(cl);rid=ref(cl,'Valid','Smith, John',2025)
    url=f'/api/v09/projects/{pid}/chapters/2/save'
    for body, ids in [(f'[@ref:{rid}]',[]),('[@ref:imaginary]',[rid]),(f'[@ref:{rid}]', ['invalid'])]:
        r=cl.post(url,json={'body':body,'reference_ids':ids,'style':'ieee'})
        assert r.status_code==422,(body,r.text)
    r=cl.post(url,json={'body':f'[@ref:{rid}]','reference_ids':[rid],'style':'ieee'})
    assert r.status_code==200,r.text
    assert r.json()['citation_integrity']=='linked_ids_format_not_independently_verified'
    assert cl.get(f'/api/v09/projects/{pid}').json()['chapters']['2']['citation_integrity']=='linked_ids_format_not_independently_verified'


def test_legacy_manual_ieee_remains_unverified_no_fabricated_mapping():
    cl=owner();pid=project(cl);rid=ref(cl,'Old','Smith, John',2025)
    url=f'/api/v09/projects/{pid}/chapters/2'
    saved=cl.post(url+'/save',json={'body':'Legacy [2] text','style':'ieee','reference_ids':[rid]})
    assert saved.status_code==200
    assert saved.json()['citation_integrity']=='legacy_numeric_needs_review'
    assert cl.get(url).json()[0]['citation_integrity']=='legacy_numeric_needs_review'
    doc=cl.get(url+'/docx')
    from docx import Document
    text=' '.join(x.text for x in Document(io.BytesIO(doc.content)).paragraphs)
    assert 'CITATION WARNING' in text
    assert 'Legacy [2]' in text
