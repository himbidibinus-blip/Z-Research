"""Offline, deterministic v0.5 integration tests. No real AI or external API calls."""
import io
import numpy as np
import pandas as pd
from fastapi.testclient import TestClient
from app.main import app
from app import ai_features as ai
from app import v05

def setup_client():
    cl=TestClient(app)
    response=cl.post('/api/auth/setup',json={'username':'testperson','password':'this-is-a-long-password'})
    assert response.status_code==200,response.text
    return cl

def upload(cl):
    with open('sample/consumer_survey_demo.csv','rb') as f:
        resp=cl.post('/api/datasets',files={'file':('demo.csv',f,'text/csv')})
    assert resp.status_code==200,resp.text
    return resp.json()['id']

def model():
    nodes=[{'id':'smm','label':'Social Media Marketing','type':'latent','indicators':['social_media_1','social_media_2','social_media_3'],'measurement':'reflective','x':150,'y':250}, {'id':'tr','label':'Trust','type':'latent','indicators':['trust_1','trust_2','trust_3'],'measurement':'reflective','x':460,'y':100},{'id':'pi','label':'Purchase Intention','type':'latent','indicators':['purchase_intention'],'measurement':'reflective','x':770,'y':250}]
    edges=[{'source':'smm','target':'tr','label':'H1'},{'source':'tr','target':'pi','label':'H2'},{'source':'smm','target':'pi','label':'H3'}]
    return nodes,edges

def test_pls_sem_pipeline_and_model_constraints():
    cl=setup_client();did=upload(cl);nodes,edges=model()
    unauth=TestClient(app)
    assert unauth.post('/api/v05/pls/run',json={'dataset_id':did,'nodes':nodes,'edges':edges}).status_code==401
    q={'dataset_id':did,'nodes':nodes,'edges':edges,'bootstraps':65,'seed':7}
    response=cl.post('/api/v05/pls/run',json=q)
    assert response.status_code==200,response.text
    result=response.json()['result']
    assert result['n']==120 and result['excluded_rows']==0
    assert result['converged']
    assert result['bootstrap']['Social Media Marketing -> Trust']['ci95'] is not None
    assert 0<result['paths']['Social Media Marketing -> Trust']<1
    assert len(result['outer_model'])==7
    assert len(result['cross_loadings'])==7
    assert 'Trust' in result['reliability_validity']
    assert 'Social Media Marketing | Trust' in result['htmt']
    assert result['bootstraps_successful']==65
    again=cl.post('/api/v05/pls/run',json=q)
    assert response.json()['result']['bootstrap']==again.json()['result']['bootstrap']
    assert cl.get('/api/v05/analyses').status_code==200
    analysis_id=response.json()['analysis_id']
    prompt=cl.post('/api/v05/chapters/preview',json={'chapter':4,'dataset_id':did,'analysis_id':analysis_id,'style':'ieee','research_facts':'Survey method unknown'})
    assert prompt.status_code==200,prompt.text
    assert '0.' in prompt.json()['prompt']
    assert 'never invent' in prompt.json()['prompt'].lower()
    invalid=cl.post('/api/v05/pls/run',json={'dataset_id':did,'nodes':nodes,'edges':[{'source':'smm','target':'tr'},{'source':'tr','target':'smm'}],'bootstraps':0})
    assert invalid.status_code==422
    assert 'Cycles' in invalid.json()['detail']
    unknown=cl.post('/api/v05/pls/run',json={'dataset_id':did,'nodes':[dict(nodes[0],indicators=['not_a_column']),nodes[1]],'edges':[edges[0]],'bootstraps':0})
    assert unknown.status_code==422


def test_references_word_and_guides(monkeypatch):
    cl=setup_client();did=upload(cl)
    a=cl.post('/api/v05/references/manual',json={'title':'An article','authors':'Researcher A','year':2024,'doi':'https://doi.org/10.1234/example','venue':'Sample Journal','issn':'1234-5678'})
    assert a.status_code==200,a.text
    assert a.json()['scopus_status']=='NOT_VERIFIED'
    assert a.json()['doi']=='10.1234/example'
    assert len(cl.get('/api/v05/references').json())==1
    uploadr=cl.post('/api/v05/references/upload',files={'file':('notes.txt',b'Writing guide for chapter 3 and chapter 4. DOI: 10.2345/something','text/plain')})
    assert uploadr.status_code==200,uploadr.text
    assert uploadr.json()['possible_doi']=='10.2345/something'
    gid=uploadr.json()['id']
    prompt=cl.post('/api/v05/chapters/preview',json={'chapter':3,'dataset_id':did,'guide_id':gid,'style':'apa','research_facts':'Sampling not recorded.'})
    assert prompt.status_code==200,prompt.text
    assert 'DOI: 10.2345/something' in prompt.json()['prompt']
    export=cl.post('/api/v05/chapters/export-word',json={'chapter':3,'body':'# 3.1 Study Design\nThis is a **test** with no invented sampling.\n\n| Variable | Value |\n|---|---|\n| Alpha | 0.8 |'})
    assert export.status_code==200,export.text
    from docx import Document
    doc=Document(io.BytesIO(export.content))
    assert len(doc.tables)==1
    assert '3.1 Study Design' in '\n'.join(p.text for p in doc.paragraphs)
    assert cl.delete('/api/v05/references/'+a.json()['id']).status_code==200
    assert cl.get('/api/v05/references').json()==[]


def test_chapter_ai_provider_mock(monkeypatch):
    cl=setup_client();did=upload(cl)
    payload={'chapter':4,'dataset_id':did,'research_facts':'No actual sampling information available','style':'ieee'}
    assert cl.post('/api/v05/chapters/generate',json=payload).status_code==403
    sid=cl.cookies['qr_session']
    ai.SESSIONS[sid]['api_key']='offline-mock-only'
    ai.SESSIONS[sid]['model']='simulated-model'
    async def fake_ai(key,model,system,user_prompt,max_tokens=1800):
        assert max_tokens==6500
        assert 'NO COMPLETED ANALYSIS PROVIDED' in user_prompt
        assert 'fabricate' in system.lower()
        return '# 4.1 Findings\n[RUN ANALYSIS] No results claimed.'
    monkeypatch.setattr(v05,'openai_generate',fake_ai)
    resp=cl.post('/api/v05/chapters/generate',json=payload)
    assert resp.status_code==200,resp.text
    ch_id=resp.json()['id']
    assert cl.get(f'/api/v05/chapters/{ch_id}').status_code==200
    doc=cl.get(f'/api/v05/chapters/{ch_id}/docx')
    assert doc.status_code==200
    assert doc.content[:2]==b'PK'
