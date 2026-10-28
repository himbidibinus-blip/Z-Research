"""Z-Research v0.9 connected project workflow.
Local-only, owner authenticated. Legacy /api/v05 contracts remain available.
Snapshots prevent reporting stale results after dataset/model edits.
"""
from __future__ import annotations
import hashlib, io, json, re, sqlite3, uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Literal, Optional
import numpy as np
import pandas as pd
from fastapi import APIRouter, Cookie, HTTPException, Request
from fastapi.responses import Response
from pydantic import BaseModel, Field
from .ai_features import auth, db, now
from .sem_engine import SEMError, estimate, validate_model
from .providers import generate_text
from .v05 import ChapterRequest, chapter_prompt, GUARD
from .citations import cited_ids, validate_linked_sources, render_document, TOKEN, BIBLIO, LEGACY_NUMERIC

router = APIRouter(prefix='/api/v09',tags=['v0.9 Connected Research Workspace'])
DATA = Path(__file__).resolve().parents[1] / 'local_data'


def initialize():
 with db() as c:
  c.execute('''CREATE TABLE IF NOT EXISTS v09_projects (
   id TEXT PRIMARY KEY,user_id TEXT NOT NULL,name TEXT NOT NULL,dataset_id TEXT,
   model_json TEXT NOT NULL DEFAULT '{"nodes":[],"edges":[]}',revision INTEGER NOT NULL DEFAULT 1,
   created_at TEXT NOT NULL,updated_at TEXT NOT NULL)''')
  c.execute('''CREATE TABLE IF NOT EXISTS v09_events (
   id INTEGER PRIMARY KEY AUTOINCREMENT,project_id TEXT NOT NULL,event_type TEXT NOT NULL,
   details TEXT NOT NULL,created_at TEXT NOT NULL)''')
  c.execute('''CREATE TABLE IF NOT EXISTS v09_runs (
   id TEXT PRIMARY KEY,project_id TEXT NOT NULL,analysis_id TEXT NOT NULL,
   fingerprint TEXT NOT NULL,quality_json TEXT NOT NULL,created_at TEXT NOT NULL)''')
  c.execute('''CREATE TABLE IF NOT EXISTS v09_chapters (
   id TEXT PRIMARY KEY,project_id TEXT NOT NULL,user_id TEXT NOT NULL,chapter INTEGER NOT NULL,
   body TEXT NOT NULL,source_fingerprint TEXT NOT NULL,analysis_id TEXT,
   reference_ids TEXT NOT NULL,created_at TEXT NOT NULL,updated_at TEXT NOT NULL)''')
  # Compatible migration: run on startup and on clean test databases.
  cols={r['name'] for r in c.execute('PRAGMA table_info(v09_chapters)').fetchall()}
  if 'evidence_fingerprint' not in cols:
   c.execute('ALTER TABLE v09_chapters ADD COLUMN evidence_fingerprint TEXT')
  if 'citation_style' not in cols:
   c.execute("ALTER TABLE v09_chapters ADD COLUMN citation_style TEXT NOT NULL DEFAULT 'apa'")
initialize()


def _snapshot_rows(connection, uid, ids):
    placeholders = ','.join('?' for _ in ids)
    return connection.execute(
        f'SELECT id,title,authors,year,venue,doi,url,scopus_status FROM research_references WHERE user_id=? AND id IN ({placeholders})',
        (uid,*ids)).fetchall()


# Heuristic detector, NOT a source verifier. Deliberately conservative:
# parenthetical and narrative author-year forms trigger a manual review alert.
# Linked [@ref:<id>] tokens remain authoritative for ID integrity.
PAREN_AUTHOR_YEAR = re.compile(
 r'\((?:[A-ZÀ-ÖØ-Þ][\wÀ-ž\-’\' ]{1,90}(?:\s*(?:&|and|et al\.)\s*[A-ZÀ-ÖØ-Þ]?[\wÀ-ž\-’\' ]{0,80})?\s*,\s*)'
 r'(?:19|20)\d{2}[a-z]?\s*\)',re.I)
NARRATIVE_AUTHOR_YEAR = re.compile(
 r'\b[A-ZÀ-ÖØ-Þ][A-Za-zÀ-ž\-’\']{2,60}(?:\s+(?:et al\.|&\s+[A-Z][A-Za-z-]+))?\s*\(\s*(?:19|20)\d{2}[a-z]?\s*\)')


def likely_unlinked_citations(body):
 """Return examples of *possible* legacy citations; never assert factual truth."""
 text=TOKEN.sub('',body or '')
 candidates=[]
 for pattern in (PAREN_AUTHOR_YEAR,NARRATIVE_AUTHOR_YEAR):
  for match in pattern.finditer(text):
   value=match.group().strip()
   if value not in candidates:candidates.append(value)
 return candidates[:8]


def chapter_citation_integrity(body, selected):
    linked = cited_ids(body)
    plain_candidates = likely_unlinked_citations(body)
    missing = sorted(set(linked)-set(selected))
    if missing:
        return {'citation_integrity':'broken_citation', 'unselected_citation_ids':missing, 'unlinked_author_year_candidates':plain_candidates}
    if linked:
        if LEGACY_NUMERIC.search(TOKEN.sub('',body or '')):
            return {'citation_integrity':'mixed_linked_and_legacy_needs_review', 'unselected_citation_ids':[], 'unlinked_author_year_candidates':plain_candidates}
        return {'citation_integrity':'linked_ids_format_not_independently_verified', 'unselected_citation_ids':[], 'unlinked_author_year_candidates':plain_candidates}
    if LEGACY_NUMERIC.search(body or ''):
        return {'citation_integrity':'legacy_numeric_needs_review', 'unselected_citation_ids':[], 'unlinked_author_year_candidates':plain_candidates}
    # Plain prose, AI-generated citations, and author-year strings cannot be
    # certified solely by matching selected bibliographic IDs.
    return {'citation_integrity':'unverified_plain_text', 'unselected_citation_ids':[], 'unlinked_author_year_candidates':plain_candidates}

def reference_snapshot(uid, reference_ids, connection=None):
    """Validate reference ownership and snapshot bibliographic content for a chapter."""
    if len(reference_ids)!=len(set(reference_ids)):
        raise HTTPException(422,'Duplicate reference ID in selected sources')
    if not reference_ids:
        return hashlib.sha256(b'[]').hexdigest()
    if connection is None:
        with db() as c:
            rows=_snapshot_rows(c,uid,reference_ids)
    else:
        rows=_snapshot_rows(connection,uid,reference_ids)
    if len(rows)!=len(reference_ids):
        raise HTTPException(422,'Selected reference ID does not exist in your library')
    keyed={r['id']:dict(r) for r in rows}
    normalized=[keyed[rid] for rid in reference_ids]
    return hashlib.sha256(json.dumps(normalized,sort_keys=True,ensure_ascii=False,default=str).encode('utf-8')).hexdigest()


def chapter_evidence(row, uid, data_fp):
    selected=json.loads(row['reference_ids'])
    with db() as c:
        if selected:
            refs=c.execute('SELECT id FROM research_references WHERE user_id=? AND id IN ('+','.join('?' for _ in selected)+')',(uid,*selected)).fetchall()
            found={r['id'] for r in refs}
        else:found=set()
    broken=sorted(set(selected)-found)
    if len(selected)!=len(set(selected)):
        broken=sorted(set(broken)|set(selected))
    saved_evidence=row['evidence_fingerprint'] if 'evidence_fingerprint' in row.keys() else None
    if broken:status='broken_citation'
    elif row['source_fingerprint']!=data_fp:status='stale'
    elif saved_evidence:
        current=reference_snapshot(uid,selected)
        status='current' if current==saved_evidence else 'stale'
    elif selected:status='needs_review'  # prior-version chapter lacked evidence snapshot
    else:status='current'
    linked=chapter_citation_integrity(row['body'] if 'body' in row.keys() else '',selected)
    if linked['unselected_citation_ids']:
        status='broken_citation'
        broken=sorted(set(broken)|set(linked['unselected_citation_ids']))
    elif linked['citation_integrity'] in ('legacy_numeric_needs_review','mixed_linked_and_legacy_needs_review') and status=='current':
        status='needs_review'
    return {'status':status,'broken_reference_ids':broken,
            'evidence_fingerprint':saved_evidence,**linked}


def validated_chapter_sources(uid, ids):
    return reference_snapshot(uid,ids)


def authorized(request, cookie):return auth(request,cookie)['user_id']

def project_owned(pid,uid):
 with db() as c:row=c.execute('SELECT * FROM v09_projects WHERE id=? AND user_id=?',(pid,uid)).fetchone()
 if not row:raise HTTPException(404,'Research project not found')
 result=dict(row);result['model']=json.loads(result.pop('model_json'));return result

def active_dataset(project):
 did=project['dataset_id']
 if not did:raise HTTPException(422,'Select a dataset for this project first')
 from .main import read_df
 return read_df(did)

def fingerprint(project):
 """Digest reflects actual dataset bytes and current model; not only DB timestamp."""
 model=project['model']
 ds=project['dataset_id']
 h=hashlib.sha256()
 h.update(json.dumps(model,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode())
 h.update(b'\0');h.update((ds or '').encode())
 if ds:
  filename=DATA/(ds+'.csv')
  if not filename.is_file():return 'MISSING_DATASET'
  with filename.open('rb') as f:
   for chunk in iter(lambda:f.read(1024*1024),b''):h.update(chunk)
 return h.hexdigest()

def stamp_event(pid,kind,details):
 with db() as c:c.execute('INSERT INTO v09_events (project_id,event_type,details,created_at) VALUES (?,?,?,?)',(pid,kind,json.dumps(details,ensure_ascii=False),now()))

def current_run(pid,fprint):
 with db() as c:
  row=c.execute('SELECT * FROM v09_runs WHERE project_id=? ORDER BY created_at DESC LIMIT 1',(pid,)).fetchone()
 if not row:return None
 item=dict(row);item['quality']=json.loads(item.pop('quality_json'));item['status']='current' if item['fingerprint']==fprint else 'stale';return item

def latest_chapters(pid):
 with db() as c:rows=c.execute('SELECT id,chapter,body,substr(body,1,160) excerpt,source_fingerprint,analysis_id,reference_ids,evidence_fingerprint,citation_style,created_at,updated_at FROM v09_chapters WHERE project_id=? ORDER BY updated_at DESC, rowid DESC',(pid,)).fetchall()
 result={}
 for x in rows:
  item=dict(x)
  if item['chapter'] not in result:result[item['chapter']]=item
 return result

class ProjectCreate(BaseModel):
 name:str=Field(min_length=3,max_length=140)
 dataset_id:Optional[str]=None

@router.post('/projects')
def create_project(payload:ProjectCreate,request:Request,qr_session:Optional[str]=Cookie(None)):
 uid=authorized(request,qr_session)
 if payload.dataset_id:
  from .main import read_df
  read_df(payload.dataset_id)
 pid=uuid.uuid4().hex
 with db() as c:c.execute('INSERT INTO v09_projects VALUES (?,?,?,?,?,?,?,?)',(pid,uid,payload.name,payload.dataset_id,json.dumps({'nodes':[],'edges':[]}),1,now(),now()))
 stamp_event(pid,'created',{'name':payload.name})
 return project_owned(pid,uid)

@router.get('/projects')
def list_projects(request:Request,qr_session:Optional[str]=Cookie(None)):
 uid=authorized(request,qr_session)
 with db() as c: rows=c.execute('SELECT id,name,dataset_id,revision,updated_at FROM v09_projects WHERE user_id=? ORDER BY updated_at DESC, rowid DESC LIMIT 100',(uid,)).fetchall()
 return [dict(x) for x in rows]

@router.get('/projects/{pid}')
def get_project(pid:str,request:Request,qr_session:Optional[str]=Cookie(None)):
 p=project_owned(pid,authorized(request,qr_session));f=fingerprint(p);run=current_run(pid,f);chapters=latest_chapters(pid)
 return {**p,'fingerprint':f,'analysis':run,'chapters':{str(n):{**x,**chapter_evidence(x,p['user_id'],f)} for n,x in chapters.items()}}

class ProjectUpdate(BaseModel):
 name:Optional[str]=Field(default=None,min_length=3,max_length=140)
 dataset_id:Optional[str]=None
 model:Optional[dict]=None

@router.patch('/projects/{pid}')
def update_project(pid:str,payload:ProjectUpdate,request:Request,qr_session:Optional[str]=Cookie(None)):
 uid=authorized(request,qr_session);p=project_owned(pid,uid)
 updates=payload.model_fields_set
 if 'dataset_id' in updates and payload.dataset_id:
  from .main import read_df
  read_df(payload.dataset_id)
 name=payload.name if 'name' in updates else p['name']
 did=payload.dataset_id if 'dataset_id' in updates else p['dataset_id']
 model=payload.model if 'model' in updates else p['model']
 if not isinstance(model,dict) or not isinstance(model.get('nodes'),list) or not isinstance(model.get('edges'),list):raise HTTPException(422,'Model must have nodes and edges arrays')
 if len(model['nodes'])>24 or len(model['edges'])>100 or len(json.dumps(model))>75000:raise HTTPException(422,'Model too large')
 ids=[n.get('id') for n in model['nodes']]
 if not all(isinstance(i,str) and 0<len(i)<80 for i in ids) or len(set(ids))!=len(ids):raise HTTPException(422,'Unique construct IDs required')
 if any(e.get('source') not in ids or e.get('target') not in ids for e in model['edges']):raise HTTPException(422,'Hypothesis references unknown construct')
 if len(model['nodes'])>=2 and model['edges'] and did:
  from .main import read_df
  df,_=read_df(did)
  try:validate_model(df,model['nodes'],model['edges'])
  except SEMError as e:raise HTTPException(422,str(e)) from e
 with db() as c:c.execute('UPDATE v09_projects SET name=?,dataset_id=?,model_json=?,revision=revision+1,updated_at=? WHERE id=? AND user_id=?',(name,did,json.dumps(model,ensure_ascii=False),now(),pid,uid))
 stamp_event(pid,'edited',{'fields':list(updates)})
 return get_project(pid,request,qr_session)

def infer_constructs(df):
 groups={}
 for col in df.select_dtypes(include='number').columns:
  if re.fullmatch(r'(?i)(id|response_?id|timestamp|row_?id|age|year|gender|respondent_?id)',str(col)):continue
  match=re.fullmatch(r'(.+?)[_\- ](?:item[_\- ]?)?(\d{1,3})',str(col),flags=re.I)
  group=(match.group(1) if match else str(col)).strip('_- ').replace('_',' ').replace('-',' ')
  if group:
   key=group.lower();groups.setdefault(key,{'key':key,'label':group,'indicators':[]})['indicators'].append(str(col))
 return sorted(groups.values(),key=lambda x:x['label'].casefold())

@router.get('/projects/{pid}/constructs')
def available_constructs(pid:str,request:Request,qr_session:Optional[str]=Cookie(None)):
 p=project_owned(pid,authorized(request,qr_session));df,_=active_dataset(p)
 return {'groups':infer_constructs(df),'warning':'Column-name grouping does not establish valid theoretical constructs. Confirm every indicator and construct role.'}

class AutoModel(BaseModel):
 independent:list[str]=Field(min_length=1,max_length=12)
 dependent:str
 mediators:list[str]=Field(default_factory=list,max_length=8)
 moderators:list[str]=Field(default_factory=list,max_length=8)

@router.post('/projects/{pid}/auto-model')
def auto_model(pid:str,body:AutoModel,request:Request,qr_session:Optional[str]=Cookie(None)):
 p=project_owned(pid,authorized(request,qr_session));df,_=active_dataset(p)
 groups={g['key']:g for g in infer_constructs(df)}
 ordered=list(dict.fromkeys(body.independent+body.mediators+body.moderators+[body.dependent]))
 if len(ordered)!=len(body.independent+body.mediators+body.moderators+[body.dependent]):raise HTTPException(422,'Construct roles must be distinct')
 if len(ordered)>24 or any(k not in groups for k in ordered):raise HTTPException(422,'Select valid, distinct constructs from the dataset')
 ids={k:'v_'+str(i+1) for i,k in enumerate(ordered)}
 nodes=[]
 for i,k in enumerate(ordered):
  role='dependent' if k==body.dependent else 'mediator' if k in body.mediators else 'moderator' if k in body.moderators else 'independent'
  nodes.append({'id':ids[k],'label':groups[k]['label'],'type':'latent','measurement':'reflective','role':role,'indicators':groups[k]['indicators'],'x':180 if role=='independent' else 770 if role=='dependent' else 465,'y':115+i*48})
 edges=[];used=set()
 def edge(a,b,kind='direct',moderator=None):
  key=(a,b,kind,moderator)
  if a==b or key in used:return
  used.add(key)
  rec={'source':ids[a],'target':ids[b],'kind':kind,'label':'H'+str(len(edges)+1)}
  if moderator:rec['moderator']=ids[moderator]
  edges.append(rec)
 for x in body.independent:
  edge(x,body.dependent)
  for med in body.mediators:edge(x,med)
 for med in body.mediators:edge(med,body.dependent)
 for mod in body.moderators:
  edge(mod,body.dependent)
  for x in body.independent:edge(x,body.dependent,'moderation',mod)
 model={'nodes':nodes,'edges':edges}
 update_project(pid,ProjectUpdate(model=model),request,qr_session)
 return {'model':model,'note':'Proposed hypotheses, not verified theory. Edit or remove any paths that lack theoretical justification.'}

def model_quality(out,df,model):
 alerts=[]
 if out['excluded_rows']>0:alerts.append({'severity':'review','text':f"{out['excluded_rows']} incomplete rows excluded by listwise deletion."})
 if out['bootstraps_requested']<1000:alerts.append({'severity':'review','text':'Fewer than 1,000 bootstrap samples: suitable for preview, not a definitive inferential report.'})
 ratio=(out['bootstraps_successful']/out['bootstraps_requested']) if out['bootstraps_requested'] else None
 if ratio is not None and ratio<0.95:alerts.append({'severity':'warning','text':'Bootstrap convergence success below 95%; review data and model.'})
 for k,v in out.get('reliability_validity',{}).items():
  if v.get('alpha') is not None and v['alpha']<.70:alerts.append({'severity':'review','text':f'{k}: alpha below .70 (context-dependent heuristic).'})
  if v.get('ave') is not None and v['ave']<.50:alerts.append({'severity':'review','text':f'{k}: AVE below .50 (reflective model heuristic).'})
 for k,v in out.get('htmt',{}).items():
  if v is not None and v>=.9:alerts.append({'severity':'review','text':f'HTMT for {k} at or above .90; examine discriminant validity.'})
 for k,v in out.get('inner_vif',{}).items():
  if v is not None and v>5:alerts.append({'severity':'review','text':f'{k}: high predictor collinearity (VIF > 5).'})
 alerts.append({'severity':'info','text':'Experimental centroid-weighting PLS estimator: comparison with SmartPLS/cSEM is required for publication.'})
 return {'alerts':alerts,'sample_n':out['n'],'excluded_n':out['excluded_rows'],'bootstrap_success_rate':ratio,'method_status':'EXPERIMENTAL_NOT_BENCHMARK_VALIDATED'}

class AnalyzeSettings(BaseModel):
 bootstraps:int=Field(default=200,ge=0,le=2000)
 seed:int=Field(default=42,ge=0,le=2147483647)

@router.post('/projects/{pid}/analyze')
def analyze_project(pid:str,body:AnalyzeSettings,request:Request,qr_session:Optional[str]=Cookie(None)):
 p=project_owned(pid,authorized(request,qr_session));df,_=active_dataset(p)
 model=p['model'];nodes=model['nodes'];edges=model['edges']
 if len(nodes)<2 or not edges:raise HTTPException(422,'Create a structural research model before analysis')
 fp=fingerprint(p)
 try:out=estimate(df,nodes,edges,body.bootstraps,body.seed)
 except (SEMError,ValueError,np.linalg.LinAlgError) as e:raise HTTPException(422,str(e)) from e
 quality=model_quality(out,df,model)
 descriptives=[]
 for node in nodes:
  for col in node['indicators']:
   s=pd.to_numeric(df[col],errors='coerce').dropna()
   descriptives.append({'construct':node['label'],'indicator':col,'n':int(len(s)),'mean':float(s.mean()) if len(s) else None,'sd':float(s.std(ddof=1)) if len(s)>1 else None,'missing':int(len(df)-len(s))})
 out['indicator_descriptives']=descriptives
 out['quality_review']=quality
 from .main import save_result
 save=save_result(p['dataset_id'],'pls_sem',{'model':model,'bootstraps':body.bootstraps,'seed':body.seed,'snapshot_sha256':fp},out)
 rid=uuid.uuid4().hex
 with db() as c:c.execute('INSERT INTO v09_runs VALUES (?,?,?,?,?,?)',(rid,pid,save['analysis_id'],fp,json.dumps(quality),now()))
 stamp_event(pid,'analysis_run',{'analysis_id':save['analysis_id'],'sha256':fp,'seed':body.seed,'bootstraps':body.bootstraps})
 return {'id':rid,'analysis_id':save['analysis_id'],'fingerprint':fp,'status':'current','result':save['result'],'quality':quality}

@router.get('/projects/{pid}/analysis')
def analysis_report(pid:str,request:Request,qr_session:Optional[str]=Cookie(None)):
 p=project_owned(pid,authorized(request,qr_session));run=current_run(pid,fingerprint(p))
 if not run:raise HTTPException(404,'No analysis yet. Calculate SEM-PLS first.')
 from .main import get_analysis
 item=get_analysis(run['analysis_id'])
 return {'run':run,'analysis':item,'warning':'Stale results must not be used for Chapter 4 until analysis is recalculated.' if run['status']=='stale' else 'Experimental estimator; independently verify results before publication.'}

@router.get('/projects/{pid}/evidence')
def evidence_matrix(pid:str,request:Request,qr_session:Optional[str]=Cookie(None)):
 p=project_owned(pid,authorized(request,qr_session));uid=p['user_id']
 with db() as c:
  rows=c.execute('SELECT id,title,authors,year,venue,doi,scopus_status FROM research_references WHERE user_id=? ORDER BY year DESC LIMIT 300',(uid,)).fetchall()
  news=c.execute('SELECT id,title,url,source,published FROM research_news WHERE user_id=? ORDER BY created_at DESC LIMIT 100',(uid,)).fetchall()
 refs=[dict(x) for x in rows];themes=[n.get('label','') for n in p['model'].get('nodes',[])]
 matrix=[]
 for theme in themes:
  matches=[r for r in refs if theme.casefold() in r['title'].casefold()]
  matrix.append({'construct':theme,'title_matches':len(matches),'candidate_references':matches[:12], 'status':'TITLE_MATCH_ONLY_NOT_EVIDENCE_OF_FINDINGS'})
 return {'construct_matrix':matrix,'references':refs,'news_leads':[dict(x) for x in news],'research_gap_prompts':[f'Which mechanisms and boundary conditions explain {e.get("label","a proposed relationship")}: {next((n.get("label") for n in p["model"]["nodes"] if n["id"]==e.get("source")),"X")} → {next((n.get("label") for n in p["model"]["nodes"] if n["id"]==e.get("target")),"Y")}? ' for e in p['model'].get('edges',[])][:20], 'notice':'Literature matching is by title only; findings, theory fit, Scopus indexing and novelty must be verified using original sources.'}

class DraftRequest09(BaseModel):
 research_facts:str=Field(default='',max_length=16000)
 instructions:str=Field(default='',max_length=6000)
 reference_ids:list[str]=Field(default_factory=list,max_length=120)
 style:Literal['apa','ieee','generic','custom']='apa'
 language:Literal['English','Indonesian']='English'
 custom_style:str=Field(default='',max_length=5000)
 mode:Literal['preview','generate']='preview'

@router.post('/projects/{pid}/chapters/{chapter}/draft')
async def draft_chapter(pid:str,chapter:int,body:DraftRequest09,request:Request,qr_session:Optional[str]=Cookie(None)):
 if chapter not in (1,2,3,4):raise HTTPException(422,'Only chapters 1 to 4 are supported')
 uid=authorized(request,qr_session);p=project_owned(pid,uid);fp=fingerprint(p)
 evidence=validated_chapter_sources(uid,body.reference_ids)
 run=current_run(pid,fp)
 if chapter==4 and (not run or run['status']!='current'):
  raise HTTPException(409,'Chapter 4 requires a current, computed analysis. Recalculate SEM-PLS after any dataset/model change.')
 req=ChapterRequest(chapter=chapter,dataset_id=p['dataset_id'],analysis_id=run['analysis_id'] if run and run['status']=='current' else None,model_spec=p['model'] if p['model']['nodes'] else None,research_facts=body.research_facts,instructions=body.instructions,reference_ids=body.reference_ids,style=body.style,language=body.language,custom_style=body.custom_style)
 prompt=chapter_prompt(req,uid)
 if body.mode=='preview':return {'prompt':prompt,'system':GUARD,'source_status':'current' if run and run['status']=='current' else 'NO_CURRENT_ANALYSIS','fingerprint':fp,'note':'Preview: no AI was called or text generated.'}
 session=auth(request,qr_session)
 if not (session.get('api_key') or session.get('provider')=='ollama'):raise HTTPException(403,'Connect Ollama or an AI provider first')
 if session.get('provider')=='openai':
  from .ai_features import openai_generate
  text=await openai_generate(session['api_key'],session['model'],GUARD,prompt,max_tokens=6500)
 else:
  text=await generate_text(session['provider'],session.get('api_key'),session['model'],GUARD,prompt,max_tokens=6500,custom_url=session.get('base_url',''))
 # Concurrent model editing: mark draft stale if context changed during the AI call.
 new_fp=fingerprint(project_owned(pid,uid))
 cid=uuid.uuid4().hex
 with db() as c:c.execute('INSERT INTO v09_chapters (id,project_id,user_id,chapter,body,source_fingerprint,analysis_id,reference_ids,created_at,updated_at,evidence_fingerprint,citation_style) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)',(cid,pid,uid,chapter,text,fp,run['analysis_id'] if run and run['status']=='current' else None,json.dumps(body.reference_ids),now(),now(),evidence,body.style))
 stamp_event(pid,'chapter_generated',{'chapter':chapter,'draft_id':cid,'context_changed_during_generation':fp!=new_fp})
 return {'id':cid,'body':text,'status':'current' if fp==new_fp else 'stale','note':'AI draft; independently verify all citations, reported results and research procedures. Draft was not automatically overwritten.'}

class SaveChapter(BaseModel):
 body:str=Field(min_length=1,max_length=250000)
 reference_ids:list[str]=Field(default_factory=list,max_length=120)
 style:Literal['apa','ieee','generic','custom']='apa'

@router.post('/projects/{pid}/chapters/{chapter}/save')
def save_chapter(pid:str,chapter:int,body:SaveChapter,request:Request,qr_session:Optional[str]=Cookie(None)):
 if chapter not in (1,2,3,4):raise HTTPException(422,'Chapter must be 1–4')
 uid=authorized(request,qr_session);p=project_owned(pid,uid);fp=fingerprint(p);run=current_run(pid,fp)
 evidence=validated_chapter_sources(uid,body.reference_ids)
 unlinked=validate_linked_sources(body.body,body.reference_ids)
 if unlinked:raise HTTPException(422,'Citation tokens refer to unselected sources: '+', '.join(unlinked[:8]))
 cid=uuid.uuid4().hex
 with db() as c:c.execute('INSERT INTO v09_chapters (id,project_id,user_id,chapter,body,source_fingerprint,analysis_id,reference_ids,created_at,updated_at,evidence_fingerprint,citation_style) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)',(cid,pid,uid,chapter,body.body,fp,run['analysis_id'] if run and run['status']=='current' else None,json.dumps(body.reference_ids),now(),now(),evidence,body.style))
 stamp_event(pid,'chapter_saved',{'chapter':chapter,'draft_id':cid})
 return {'id':cid,'status':'current',**chapter_citation_integrity(body.body,body.reference_ids),'notice':'Saved a new version. Previous versions remain in history; linked IDs do not prove academic citation accuracy.'}

@router.post('/projects/{pid}/chapters/{chapter}/versions/{revision_id}/restore')
def restore_chapter_version(pid:str,chapter:int,revision_id:str,request:Request,qr_session:Optional[str]=Cookie(None)):
    """Atomically restore body, selected references, and citation style as a NEW revision.

    Reuses the historical source fingerprint so changes after that historical
    version remain stale. Removed references block restoration without mutation.
    """
    if chapter not in (1,2,3,4):
        raise HTTPException(422,'Chapter must be 1–4')
    uid=authorized(request,qr_session)
    project=project_owned(pid,uid)
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        old=c.execute('SELECT * FROM v09_chapters WHERE id=? AND chapter=? AND project_id=? AND user_id=?',
                      (revision_id,chapter,pid,uid)).fetchone()
        if not old:
            raise HTTPException(404,'Chapter version not found')
        ref_ids=json.loads(old['reference_ids'])
        # Same connection and write lock: no concurrent reference removal while
        # we check and write the new version.
        current_evidence=reference_snapshot(uid,ref_ids,connection=c)
        missing=validate_linked_sources(old['body'],ref_ids)
        if missing:
            raise HTTPException(409,'Historical chapter contains unlinked citation IDs. Review before restoring: '+', '.join(missing[:8]))
        new_id=uuid.uuid4().hex
        timestamp=now()
        c.execute("""INSERT INTO v09_chapters (id,project_id,user_id,chapter,body,source_fingerprint,analysis_id,
           reference_ids,created_at,updated_at,evidence_fingerprint,citation_style)
           VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
           (new_id,pid,uid,chapter,old['body'],old['source_fingerprint'],old['analysis_id'],
            old['reference_ids'],timestamp,timestamp,old['evidence_fingerprint'],old['citation_style']))
        c.execute('INSERT INTO v09_events (project_id,event_type,details,created_at) VALUES (?,?,?,?)',
                  (pid,'chapter_restored',json.dumps({'chapter':chapter,'from_version_id':revision_id,'new_version_id':new_id}),timestamp))
        row=c.execute('SELECT * FROM v09_chapters WHERE id=?',(new_id,)).fetchone()
    result=chapter_evidence(row,uid,fingerprint(project))
    return {'id':new_id,'body':old['body'],'reference_ids':ref_ids,'citation_style':old['citation_style'],
            **result,'warning':('Restored historical evidence differs from the current project or reference library. Review and revalidate before submission.' if result['status']!='current' else '')}

@router.get('/projects/{pid}/chapters/{chapter}')
def get_chapter(pid:str,chapter:int,request:Request,qr_session:Optional[str]=Cookie(None)):
 p=project_owned(pid,authorized(request,qr_session))
 with db() as c: rows=c.execute('SELECT * FROM v09_chapters WHERE project_id=? AND chapter=? ORDER BY updated_at DESC, rowid DESC LIMIT 30',(pid,chapter)).fetchall()
 items=[];fp=fingerprint(p)
 for row in rows:
  x=dict(row);x.update(chapter_evidence(row,p['user_id'],fp));x['reference_ids']=json.loads(x['reference_ids']);items.append(x)
 return items

@router.get('/projects/{pid}/audit')
def audit(pid:str,request:Request,qr_session:Optional[str]=Cookie(None)):
 p=project_owned(pid,authorized(request,qr_session));fp=fingerprint(p);run=current_run(pid,fp);chapters=latest_chapters(pid)
 alerts=[]
 if not p['dataset_id']:alerts.append({'severity':'blocker','message':'No dataset selected'})
 if not p['model']['edges']:alerts.append({'severity':'blocker','message':'No structural hypotheses available'})
 if not run:alerts.append({'severity':'blocker','message':'No SEM-PLS calculation saved'})
 elif run['status']=='stale':alerts.append({'severity':'blocker','message':'Analysis is stale; source data or model was modified'})
 else:alerts.extend({'severity':i['severity'],'message':i['text']} for i in run['quality']['alerts'])
 for c,x in chapters.items():
  evidence=chapter_evidence(x,p['user_id'],fp)
  if evidence['status']!='current':alerts.append({'severity':'blocker' if evidence['status']=='broken_citation' else 'review','message':f'Chapter {c}: {evidence["status"]}' + (f' (missing reference IDs: {", ".join(evidence["broken_reference_ids"])})' if evidence['broken_reference_ids'] else '')})
  if evidence['unlinked_author_year_candidates']:
   examples=', '.join(evidence['unlinked_author_year_candidates'][:3])
   alerts.append({'severity':'review','message':f'Chapter {c}: possible unlinked author-year citation(s) require manual review: {examples}. Use Cite+ to link sources; this is not a factual verification.'})
  elif evidence['citation_integrity']=='unverified_plain_text' and (x.get('reference_ids') and x['reference_ids']!='[]'):
   alerts.append({'severity':'review','message':f'Chapter {c}: references selected but no linked citation tokens. Manually written citations are not verified.'})
 with db() as c:events=c.execute('SELECT event_type,details,created_at FROM v09_events WHERE project_id=? ORDER BY id DESC LIMIT 40',(pid,)).fetchall()
 return {'project':p['name'],'dataset_id':p['dataset_id'],'model_revision':p['revision'],'sha256':fp,'analysis_status':run['status'] if run else 'missing','draft_status':{str(k):chapter_evidence(x,p['user_id'],fp)['status'] for k,x in chapters.items()},'citation_integrity':{str(k):chapter_evidence(x,p['user_id'],fp)['citation_integrity'] for k,x in chapters.items()},'alerts':alerts,'events':[dict(x) for x in events],'scope':'Audit checks data/model traceability, not full academic or statistical validity.'}

@router.get('/projects/{pid}/reproducibility.json')
def reproduce(pid:str,request:Request,qr_session:Optional[str]=Cookie(None)):
 p=project_owned(pid,authorized(request,qr_session));fp=fingerprint(p);run=current_run(pid,fp)
 analysis=None
 if run:
  from .main import get_analysis
  analysis=get_analysis(run['analysis_id'])
 report={'schema':'zresearch-reproducibility-v0.9','created_at':now(),'project':p['name'],'dataset_id':p['dataset_id'],'model':p['model'],'fingerprint_sha256':fp,'analysis_status':run['status'] if run else 'missing','analysis':analysis,'experimental_disclosure':'PLS algorithm is not independently benchmarked as SmartPLS-equivalent. Dataset not included in JSON.','chapter_status':{str(n):chapter_evidence(x,p['user_id'],fp)['status'] for n,x in latest_chapters(pid).items()}}
 return Response(json.dumps(report,ensure_ascii=False,indent=2),media_type='application/json',headers={'Content-Disposition':'attachment; filename="zresearch_reproducibility_v09.json"'})

@router.post('/projects/{pid}/preview')
def preview_analysis(pid:str,request:Request,qr_session:Optional[str]=Cookie(None)):
 """Fast live re-estimation without persistence or unsupported significance claims."""
 p=project_owned(pid,authorized(request,qr_session));df,_=active_dataset(p)
 try:
  out=estimate(df,p['model']['nodes'],p['model']['edges'],bootstraps=0,seed=42)
 except (SEMError,ValueError,np.linalg.LinAlgError) as e:raise HTTPException(422,str(e)) from e
 return {'fingerprint':fingerprint(p),'n':out['n'],'paths':out['paths'],'r_squared':out['r_squared'],'reliability_validity':out['reliability_validity'],'status':'LIVE_PREVIEW_NO_INFERENCE','notice':'Fast preview without bootstrap uncertainty. Run the complete analysis for inference.'}

@router.get('/projects/{pid}/chapters/{chapter}/docx')
def export_saved_chapter_docx(pid: str, chapter: int, request: Request, qr_session: Optional[str] = Cookie(None)):
    """Export the latest *saved* version, never any unsaved browser editor buffer."""
    if chapter not in (1, 2, 3, 4):
        raise HTTPException(422, 'Chapter must be between 1 and 4')
    uid = authorized(request, qr_session)
    project = project_owned(pid, uid)
    with db() as c:
        row = c.execute('SELECT * FROM v09_chapters WHERE project_id=? AND user_id=? AND chapter=? ORDER BY updated_at DESC, rowid DESC LIMIT 1', (pid, uid, chapter)).fetchone()
    if not row:
        raise HTTPException(404, 'Save a chapter draft before exporting to Word')
    evidence = chapter_evidence(row, uid, fingerprint(project))
    from docx import Document
    from docx.shared import Pt
    out = io.BytesIO()
    document = Document()
    style = document.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    document.add_heading(f'Chapter {chapter}', 0)
    if evidence['status'] != 'current':
        document.add_paragraph('INTEGRITY WARNING: '+evidence['status']+'; verify unavailable or edited references before submission.')
    selected=json.loads(row['reference_ids'])
    with db() as c:
        refs=[dict(r) for r in _snapshot_rows(c,uid,selected)] if selected else []
    rendered=render_document(row['body'],row['citation_style'],refs)
    if rendered['missing_reference_ids'] or rendered['legacy_numeric']:
        document.add_paragraph('CITATION WARNING: unresolved or legacy numeric citations require manual review.')
    for paragraph in rendered['body'].split('\n'):
        stripped = paragraph.strip()
        if stripped.startswith('### '):
            document.add_heading(stripped[4:], level=3)
        elif stripped.startswith('## '):
            document.add_heading(stripped[3:], level=2)
        elif stripped.startswith('# '):
            document.add_heading(stripped[2:], level=1)
        else:
            document.add_paragraph(paragraph)
    document.add_paragraph('')
    document.add_paragraph('Z-Research research-integrity notice: AI-generated language, citations, research procedures, and experimental SEM-PLS calculations require independent verification before submission.')
    document.save(out)
    return Response(
        out.getvalue(),
        media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        headers={'Content-Disposition': f'attachment; filename="Z-Research_Chapter_{chapter}_v09.docx"', 'Cache-Control': 'no-store'},
    )
