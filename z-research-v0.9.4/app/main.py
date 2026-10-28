import io, json, os, re, sqlite3, uuid, hashlib, csv, unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional
import numpy as np
import pandas as pd
from fastapi import FastAPI, File, Form, HTTPException, UploadFile, Request, Cookie, Depends
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from scipy import stats
import statsmodels.api as sm

BASE = Path(__file__).resolve().parents[1]
DATA = BASE / 'local_data'
DATA.mkdir(exist_ok=True)
DB = DATA / 'quantresearch.db'
app = FastAPI(title='Z-Research Connected Research Workspace',version='0.9.4',description='Local-first research projects, data/model provenance, experimental SEM-PLS, citation-linked chapters, and Ollama/cloud drafting')
app.mount('/static', StaticFiles(directory=str(BASE/'app'/'static')), name='static')

def conn():
    c=sqlite3.connect(DB)
    c.row_factory=sqlite3.Row
    return c

def init():
    with conn() as c:
        c.execute('CREATE TABLE IF NOT EXISTS datasets (id TEXT PRIMARY KEY, name TEXT NOT NULL, filename TEXT NOT NULL, rows INTEGER NOT NULL, columns_json TEXT NOT NULL, created_at TEXT NOT NULL)')
        c.execute('''CREATE TABLE IF NOT EXISTS dataset_originals (dataset_id TEXT PRIMARY KEY, original_name TEXT NOT NULL, stored_path TEXT NOT NULL, sha256 TEXT NOT NULL, byte_size INTEGER NOT NULL, imported_at TEXT NOT NULL, content_type TEXT NOT NULL, FOREIGN KEY(dataset_id) REFERENCES datasets(id))''')
        c.execute('CREATE TABLE IF NOT EXISTS analyses (id TEXT PRIMARY KEY, dataset_id TEXT NOT NULL, analysis_type TEXT NOT NULL, configuration TEXT NOT NULL, result TEXT NOT NULL, created_at TEXT NOT NULL, FOREIGN KEY(dataset_id) REFERENCES datasets(id))')
init()

def clean_float(x):
    if isinstance(x,dict): return {k:clean_float(v) for k,v in x.items()}
    if isinstance(x,(tuple,list)): return [clean_float(v) for v in x]
    if isinstance(x,(np.integer,)): return int(x)
    if isinstance(x,(np.floating,float)):
        n=float(x)
        return n if np.isfinite(n) else None
    if isinstance(x,(np.bool_,)): return bool(x)
    return x

def read_df(dataset_id):
    with conn() as c: row=c.execute('SELECT * FROM datasets WHERE id=?',(dataset_id,)).fetchone()
    if row is None: raise HTTPException(404,'Dataset not found')
    path=DATA/(dataset_id+'.csv')
    if not path.exists(): raise HTTPException(404,'Dataset file missing')
    return pd.read_csv(path),dict(row)

def numeric_frame(df, cols, min_n=3):
    if not cols or len(set(cols))!=len(cols): raise HTTPException(422,'Select distinct column names')
    if any(x not in df.columns for x in cols): raise HTTPException(422,'Unknown column')
    d=df[cols].apply(pd.to_numeric,errors='coerce').dropna()
    if len(d)<min_n: raise HTTPException(422,f'At least {min_n} complete numeric observations are required')
    return d

def save_result(dataset_id, kind, config, result):
    aid=uuid.uuid4().hex
    payload=clean_float(result)
    with conn() as c:
        c.execute('INSERT INTO analyses VALUES (?,?,?,?,?,?)',(aid,dataset_id,kind,json.dumps(config),json.dumps(payload,allow_nan=False),datetime.now(timezone.utc).isoformat()))
    return {'analysis_id':aid,'result':payload}

def require_owner(request:Request,qr_session:Optional[str]=Cookie(None)):
    # Dependency on every legacy dataset / statistical endpoint (incl. export).
    from .ai_features import auth
    return auth(request, qr_session)['user_id']

@app.get('/')
def index(): return FileResponse(BASE/'app'/'static'/'index.html')

@app.get('/api/health')
def health(): return {'status':'ok','version':'0.9.4','mode':'local'}

def validate_original_headers(content:bytes, suffix:str):
    """Reject source-level duplicate labels *before* Pandas rewrites them.
    Keep the uploaded bytes untouched for provenance. UTF-8 BOM and whitespace
    or unicode-normalization collisions are treated as duplicates.
    """
    if suffix == '.csv':
        try:
            source = content.decode('utf-8-sig')
        except UnicodeDecodeError:
            source = content.decode('latin-1')
        try:
            headers = next(csv.reader(io.StringIO(source, newline=''), strict=True))
        except (csv.Error, StopIteration) as exc:
            raise HTTPException(422,'Invalid or missing CSV header row') from exc
    else:
        from openpyxl import load_workbook
        try:
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                headers = list(next(book.active.iter_rows(min_row=1, max_row=1, values_only=True)))
            finally:
                book.close()
        except Exception as exc:
            # openpyxl and Python's zip/XML parsers raise several unrelated
            # exception classes for malformed user-controlled Excel files.
            # No dataset/raw archive has been persisted before this boundary.
            raise HTTPException(422,'Invalid or missing Excel header row; upload a valid .xlsx workbook') from exc
    keys = [unicodedata.normalize('NFKC',str(v if v is not None else '')).replace('\ufeff','').strip().casefold() for v in headers]
    if not keys or any(not x for x in keys):
        raise HTTPException(422,'Blank dataset header is not allowed')
    duplicates = sorted({x for x in keys if keys.count(x)>1})
    if duplicates:
        raise HTTPException(422,'Duplicate original column headers: '+', '.join(duplicates[:12])+'. Rename them in the source file before upload.')
    return headers

@app.post('/api/datasets',dependencies=[Depends(require_owner)])
async def upload_dataset(file:UploadFile=File(...)):
    name=(file.filename or 'dataset.csv')
    suffix=Path(name).suffix.lower()
    if suffix not in ('.csv','.xlsx'): raise HTTPException(415,'Only .csv and .xlsx supported')
    content=await file.read(10_000_001)
    if len(content)>10_000_000: raise HTTPException(413,'Maximum file size is 10 MB')
    validate_original_headers(content,suffix)
    try:
        if suffix=='.csv':
            try: df=pd.read_csv(io.BytesIO(content))
            except UnicodeDecodeError: df=pd.read_csv(io.BytesIO(content),encoding='latin-1')
        else: df=pd.read_excel(io.BytesIO(content))
    except Exception as e: raise HTTPException(422,'Cannot parse dataset; check file format') from e
    if df.empty: raise HTTPException(422,'Dataset cannot be empty')
    if len(df)>100000 or len(df.columns)>300: raise HTTPException(413,'Maximum 100,000 rows and 300 columns')
    if df.columns.duplicated().any(): raise HTTPException(422,'Duplicate column names')
    df.columns=[str(c).strip() for c in df.columns]
    if len(set(df.columns))!=len(df.columns): raise HTTPException(422,'Duplicate column names after trimming')
    did=uuid.uuid4().hex
    raw_dir=DATA/'raw_uploads'
    raw_dir.mkdir(parents=True,exist_ok=True)
    raw_path=raw_dir/(did+suffix)
    # Never overwrite an existing original, even if dataset working values are edited.
    with raw_path.open('xb') as raw:
        raw.write(content)
    try:
        df.to_csv(DATA/(did+'.csv'),index=False)
    except Exception:
        raw_path.unlink(missing_ok=True)
        raise
    now=datetime.now(timezone.utc).isoformat()
    with conn() as c:
        c.execute('INSERT INTO datasets VALUES (?,?,?,?,?,?)',(did,name,did+'.csv',len(df),json.dumps(list(df.columns)),now))
        c.execute('INSERT INTO dataset_originals (dataset_id,original_name,stored_path,sha256,byte_size,imported_at,content_type) VALUES (?,?,?,?,?,?,?)',
                  (did,name,str(Path('raw_uploads')/(did+suffix)),hashlib.sha256(content).hexdigest(),len(content),now,suffix))
    return {'id':did,'name':name,'rows':len(df),'columns':list(df.columns),'created_at':now}

@app.get('/api/datasets',dependencies=[Depends(require_owner)])
def datasets():
    with conn() as c: rows=c.execute('SELECT id,name,rows,columns_json,created_at FROM datasets ORDER BY created_at DESC').fetchall()
    return [{'id':r['id'],'name':r['name'],'rows':r['rows'],'columns':json.loads(r['columns_json']),'created_at':r['created_at']} for r in rows]

@app.get('/api/datasets/{dataset_id}/provenance',dependencies=[Depends(require_owner)])
def dataset_provenance(dataset_id:str):
    with conn() as c:
        exists=c.execute('SELECT id FROM datasets WHERE id=?',(dataset_id,)).fetchone()
        if not exists:raise HTTPException(404,'Dataset not found')
        row=c.execute('SELECT * FROM dataset_originals WHERE dataset_id=?',(dataset_id,)).fetchone()
    if not row:
        return {'dataset_id':dataset_id,'archive_status':'LEGACY_ORIGINAL_NOT_AVAILABLE','sha256':None,
                'notice':'Imported before v0.9.1; a byte-for-byte original cannot be reconstructed from a normalized CSV.'}
    info=dict(row)
    raw_path=DATA/info['stored_path']
    if not raw_path.is_file():
        return {'dataset_id':dataset_id,'archive_status':'MISSING_ORIGINAL','sha256':info['sha256']}
    digest=hashlib.sha256(raw_path.read_bytes()).hexdigest()
    return {'dataset_id':dataset_id,'archive_status':'VERIFIED' if digest==info['sha256'] else 'HASH_MISMATCH',
            'sha256':info['sha256'],'byte_size':info['byte_size'],'original_name':info['original_name'],
            'format':info['content_type'],'imported_at':info['imported_at']}

@app.get('/api/datasets/{dataset_id}',dependencies=[Depends(require_owner)])
def preview(dataset_id:str):
    df,info=read_df(dataset_id)
    preview=df.head(10).replace({np.nan:None})
    return clean_float({'id':dataset_id,'name':info['name'],'rows':len(df),'columns':df.columns.tolist(),'numeric_columns':[str(c) for c in df.columns if pd.api.types.is_numeric_dtype(df[c])],'missing':df.isna().sum().to_dict(),'preview':preview.to_dict(orient='records')})

class DescriptiveRequest(BaseModel):
    dataset_id:str
    columns:list[str] = Field(min_length=1)

@app.post('/api/analyze/descriptive',dependencies=[Depends(require_owner)])
def descriptive(req:DescriptiveRequest):
    df,_=read_df(req.dataset_id)
    if any(c not in df.columns for c in req.columns): raise HTTPException(422,'Unknown columns')
    out=[]
    for c in req.columns:
        series=pd.to_numeric(df[c],errors='coerce').dropna()
        if len(series)==0: raise HTTPException(422,f'{c} has no numeric values')
        out.append({'variable':c,'n':len(series),'missing':int(len(df)-len(series)),'mean':series.mean(),'sd':series.std(ddof=1) if len(series)>1 else None,'median':series.median(),'min':series.min(),'max':series.max()})
    return save_result(req.dataset_id,'descriptive',req.model_dump(),{'statistics':out,'method':'Descriptive statistics; sample SD uses ddof=1; missing observations excluded per variable.'})

class RegressionRequest(BaseModel):
    dataset_id:str
    outcome:str
    predictors:list[str]=Field(min_length=1)

@app.post('/api/analyze/regression',dependencies=[Depends(require_owner)])
def regression(req:RegressionRequest):
    if req.outcome in req.predictors: raise HTTPException(422,'Outcome must not be a predictor')
    df,_=read_df(req.dataset_id)
    cols=[req.outcome]+req.predictors
    d=numeric_frame(df,cols,min_n=4)
    if len(d)<=len(req.predictors)+1: raise HTTPException(422,'Insufficient observations for residual degrees of freedom')
    x=sm.add_constant(d[req.predictors],has_constant='add')
    if np.linalg.matrix_rank(x.to_numpy())<x.shape[1]: raise HTTPException(422,'Singular design matrix: remove constant or collinear predictors')
    fit=sm.OLS(d[req.outcome],x).fit()
    ci=fit.conf_int(alpha=.05)
    coefficients=[{'term':k,'estimate':fit.params[k],'std_error':fit.bse[k],'t':fit.tvalues[k],'p_value':fit.pvalues[k],'ci95_low':ci.loc[k,0],'ci95_high':ci.loc[k,1]} for k in fit.params.index]
    out={'n':len(d),'excluded_rows':len(df)-len(d),'r_squared':fit.rsquared,'adj_r_squared':fit.rsquared_adj,'f_statistic':fit.fvalue,'f_p_value':fit.f_pvalue,'coefficients':coefficients,'method':'Ordinary least squares regression with listwise deletion and classical standard errors; associations are not proof of causation.'}
    return save_result(req.dataset_id,'regression',req.model_dump(),out)

class ReliabilityRequest(BaseModel):
    dataset_id:str
    indicators:list[str]=Field(min_length=2)

@app.post('/api/analyze/reliability',dependencies=[Depends(require_owner)])
def reliability(req:ReliabilityRequest):
    df,_=read_df(req.dataset_id)
    d=numeric_frame(df,req.indicators,min_n=3)
    k=d.shape[1]
    variances=d.var(axis=0,ddof=1)
    totalvar=d.sum(axis=1).var(ddof=1)
    if totalvar<=0: raise HTTPException(422,'Zero composite score variance')
    alpha=k/(k-1)*(1-variances.sum()/totalvar)
    item_total=[]
    for col in d.columns:
        other=d.drop(columns=[col]).sum(axis=1)
        r=float(d[col].corr(other)) if other.var()>0 and d[col].var()>0 else None
        item_total.append({'indicator':col,'corrected_item_total_correlation':r})
    return save_result(req.dataset_id,'reliability',req.model_dump(),{'n':len(d),'excluded_rows':len(df)-len(d),'cronbach_alpha':alpha,'items':item_total,'method':'Cronbach alpha computed from raw item variances, complete cases only. Assess dimensionality and reverse-coded items before interpreting.'})

class MethodologyRequest(BaseModel):
    dataset_id:str
    design:str=Field(min_length=3)
    population:str=Field(min_length=3)
    sampling:str=Field(min_length=3)
    collection_period:str=Field(min_length=3)
    instrument:str=Field(min_length=3)
    ethical_process:str=Field(min_length=3)
    analysis_plan:str=Field(min_length=3)

@app.post('/api/reports/methodology',dependencies=[Depends(require_owner)])
def methodology(req:MethodologyRequest):
    df,_=read_df(req.dataset_id)
    # Only documented user-supplied facts; no fabricated sampling or ethics claims.
    draft=(f'Research design. The study used {req.design}. The target population was {req.population}. '
           f'Sampling approach: {req.sampling}. Data collection period: {req.collection_period}. '
           f'Measurement instrument: {req.instrument}. Ethical procedures: {req.ethical_process}. '
           f'The uploaded dataset contains {len(df)} rows and {len(df.columns)} columns; these counts have not been verified as unique eligible participants. '
           f'Analysis plan: {req.analysis_plan}. Analytical assumptions and exclusion criteria must be documented and checked against the actual workflow.')
    return save_result(req.dataset_id,'methodology',req.model_dump(),{'draft':draft,'disclosure':'Template assembled only from researcher-provided facts and dataset dimensions; researcher verification required.'})

@app.get('/api/analyses/{analysis_id}',dependencies=[Depends(require_owner)])
def get_analysis(analysis_id:str):
    with conn() as c: r=c.execute('SELECT * FROM analyses WHERE id=?',(analysis_id,)).fetchone()
    if not r: raise HTTPException(404,'Analysis not found')
    return {'id':r['id'],'dataset_id':r['dataset_id'],'analysis_type':r['analysis_type'],'configuration':json.loads(r['configuration']),'result':json.loads(r['result']),'created_at':r['created_at']}

@app.get('/api/analyses/{analysis_id}/export',dependencies=[Depends(require_owner)])
def export(analysis_id:str):
    a=get_analysis(analysis_id)
    from fastapi.responses import Response
    return Response(json.dumps(a,indent=2,ensure_ascii=False),media_type='application/json',headers={'Content-Disposition':f'attachment; filename="analysis_{analysis_id}.json"'})

# Optional authenticated drafting and visual modeling endpoints. Existing v0.1 analysis endpoints remain intact.
from .ai_features import router as v2_router
app.include_router(v2_router)

# Research workspace v0.5 — after existing routers to retain compatibility.
from .v05 import r as v05_router
app.include_router(v05_router)
app.version='0.9.4'

# V0.9 connected project workflow; legacy routes preserved.
from .v09 import router as v09_router
app.include_router(v09_router)

@app.get('/legacy')
def legacy_workspace():
    return FileResponse(BASE/'app'/'static'/'legacy.html')
